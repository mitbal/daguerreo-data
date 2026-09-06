#!/usr/bin/env python3
"""Bounded, idempotent IDX RDF earnings detector and CSV synchronizer."""

from __future__ import annotations
import argparse
import csv
import datetime as dt
import hashlib
import json
import math
import os
import shutil
import sys
import tempfile
import time
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urljoin
from xml.etree import ElementTree

try:
    import openpyxl
except ImportError:
    openpyxl = None
try:
    from curl_cffi import requests
except ImportError:
    requests = None

ROOT = Path(__file__).resolve().parents[2]
TASK = Path(__file__).resolve().parent
CSV_DIR = ROOT / "jkse" / "income_statements"
IDX_HOME = "https://www.idx.co.id/id"
IDX_API = "https://www.idx.co.id/primary/ListedCompany/GetFinancialReport"
IDX_BASE = "https://www.idx.co.id"
CSV_HEADER = [
    "date",
    "symbol",
    "reported_currency",
    "calendar_year",
    "period",
    "revenue",
    "gross_profit",
    "net_income",
]
PERIODS = ("tw1", "tw2", "tw3", "audit")
PERIOD_NAME = dict(zip(PERIODS, ("Q1", "Q2", "Q3", "Q4")))
PREVIOUS = {"tw1": None, "tw2": "tw1", "tw3": "tw2", "audit": "tw3"}
END_MONTH = {"tw1": (3, 31), "tw2": (6, 30), "tw3": (9, 30), "audit": (12, 31)}
PARENT_NI = "laba (rugi) yang dapat diatribusikan ke entitas induk"
REVENUE_LABELS = {"penjualan dan pendapatan usaha", "penjualan dan pendapatan"}
GROSS_LABELS = {"laba bruto", "jumlah laba bruto", "laba kotor", "jumlah laba kotor"}
BANK_INTEREST_INCOME_LABELS = {
    "pendapatan bunga",
    "pendapatan pengelolaan dana oleh bank sebagai mudharib",
}
BANK_INTEREST_EXPENSE_LABELS = {
    "beban bunga",
    "hak pihak ketiga atas bagi hasil dana syirkah temporer",
}
BANK_NON_INTEREST_LABELS = {
    "pendapatan investasi",
    "pendapatan provisi dan komisi dari transaksi lainnya selain kredit",
    "pendapatan transaksi perdagangan",
    "pendapatan dividen",
    "keuntungan (kerugian) yang telah direalisasi atas instrumen derivatif",
    "penerimaan kembali aset yang telah dihapusbukukan",
    "keuntungan (kerugian) selisih kurs mata uang asing",
    "keuntungan (kerugian) pelepasan aset tetap",
    "keuntungan (kerugian) pelepasan agunan yang diambil alih",
    "pendapatan operasional lainnya",
}
BANK_OPERATING_PROFIT = "jumlah laba operasional"
BANK_PRETAX_PROFIT = "jumlah laba (rugi) sebelum pajak penghasilan"
BANK_NON_OPERATING_INCOME = {
    "pendapatan bukan operasional",
    "bagian atas laba (rugi) entitas asosiasi yang dicatat dengan menggunakan metode ekuitas",
    "bagian atas laba (rugi) entitas ventura bersama yang dicatat menggunakan metode ekuitas",
}
BANK_NON_OPERATING_EXPENSE = "beban bukan operasional"
BANK_PROVISION_INCOME_PREFIXES = (
    "pemulihan penyisihan kerugian penurunan nilai",
    "pemulihan penyisihan estimasi kerugian",
    "pembalikan (beban) estimasi kerugian",
)
BANK_PROVISION_EXPENSE_PREFIXES = (
    "pembentukan kerugian penurunan nilai",
    "pembentukan penyisihan kerugian penurunan nilai",
)
RETRY = {403, 408, 425, 429, 500, 502, 503, 504}
IMPERSONATE = "chrome124"
EXTRACTOR_VERSION = 4


def norm(v):
    return " ".join(str(v or "").strip().lower().split())


def period_date(year, period):
    m, d = END_MONTH[period]
    return dt.date(year, m, d)


def discovery_periods(only_period):
    """Return the selected filing period plus its cumulative dependency."""
    if only_period is None:
        return PERIODS
    previous = PREVIOUS[only_period]
    return (only_period,) if previous is None else (previous, only_period)


def meta_value(meta, *labels, prefixes=()):
    for label in labels:
        value = meta.get(label)
        if value not in (None, ""):
            return value
    for key, value in meta.items():
        if value not in (None, "") and any(
            key.startswith(prefix) for prefix in prefixes
        ):
            return value
    return None


def parse_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if value:
        text = str(value).strip()
        for f in (
            "%Y-%m-%d",
            "%Y-%m-%d %H:%M:%S",
            "%d/%m/%Y",
            "%d-%m-%Y",
            "%B %d, %Y",
        ):
            try:
                return dt.datetime.strptime(text, f).date()
            except ValueError:
                pass
    return None


def factor_value(v):
    s = norm(v)
    if "satuan penuh" in s or "full amount" in s:
        return 1
    if "ribuan" in s or "thousand" in s:
        return 1_000
    if "jutaan" in s or "million" in s:
        return 1_000_000
    if "miliar" in s or "billion" in s:
        return 1_000_000_000
    if not s:
        return 1
    raise ValueError(f"unknown rounding unit: {v!r}")


def numeric_value(value, multiplier):
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(value)
    ):
        return None
    scaled = value * multiplier
    if isinstance(scaled, float) and scaled.is_integer():
        return int(scaled)
    return scaled


def extract_statement(sheet, multiplier, is_bank):
    found = {}
    sources = {"revenue": [], "gross_profit": [], "net_income": []}
    bank_interest_income = 0
    bank_interest_expense = 0
    bank_non_interest = 0
    bank_provision_effect = 0
    bank_operating_profit = None
    bank_pretax_profit = None
    bank_non_operating_adjustment = 0
    bank_operating_profit_derived = False
    bank_seen = {
        "interest_income": False,
        "interest_expense": False,
        "non_interest": False,
        "provisions": False,
        "operating_profit": False,
    }

    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        raw_label = str(row[0] or "").strip()
        label = norm(raw_label)
        value = numeric_value(row[1] if len(row) > 1 else None, multiplier)
        if value is None:
            continue
        if "revenue" not in found and label in REVENUE_LABELS:
            found["revenue"] = (value, raw_label)
            sources["revenue"].append(raw_label)
        if "gross_profit" not in found and label in GROSS_LABELS:
            found["gross_profit"] = (value, raw_label)
            sources["gross_profit"].append(raw_label)
        if "net_income" not in found and label == PARENT_NI:
            found["net_income"] = (value, raw_label)
            sources["net_income"].append(raw_label)

        if not is_bank:
            continue
        if label in BANK_INTEREST_INCOME_LABELS:
            bank_interest_income += value
            bank_seen["interest_income"] = True
            sources["revenue"].append(raw_label)
        elif label in BANK_INTEREST_EXPENSE_LABELS:
            bank_interest_expense += value
            bank_seen["interest_expense"] = True
            sources["revenue"].append(raw_label)
        elif label in BANK_NON_INTEREST_LABELS:
            bank_non_interest += value
            bank_seen["non_interest"] = True
            sources["revenue"].append(raw_label)
        elif label == BANK_OPERATING_PROFIT:
            bank_operating_profit = value
            bank_seen["operating_profit"] = True
            sources["gross_profit"].append(raw_label)
        elif label == BANK_PRETAX_PROFIT:
            bank_pretax_profit = value
            sources["gross_profit"].append(raw_label)
        elif label in BANK_NON_OPERATING_INCOME:
            bank_non_operating_adjustment += value
            sources["gross_profit"].append(raw_label)
        elif label == BANK_NON_OPERATING_EXPENSE:
            bank_non_operating_adjustment -= value
            sources["gross_profit"].append(raw_label)
        elif any(
            label.startswith(prefix) for prefix in BANK_PROVISION_INCOME_PREFIXES
        ):
            bank_provision_effect += value
            bank_seen["provisions"] = True
            sources["gross_profit"].append(raw_label)
        elif any(
            label.startswith(prefix) for prefix in BANK_PROVISION_EXPENSE_PREFIXES
        ):
            bank_provision_effect -= value
            bank_seen["provisions"] = True
            sources["gross_profit"].append(raw_label)

    if is_bank and "net_income" in found:
        if bank_operating_profit is None and bank_pretax_profit is not None:
            bank_operating_profit = (
                bank_pretax_profit - bank_non_operating_adjustment
            )
            bank_seen["operating_profit"] = True
            bank_operating_profit_derived = True
        required = ("interest_income", "interest_expense", "non_interest", "operating_profit")
        missing = [name for name in required if not bank_seen[name]]
        if missing:
            raise ValueError(f"missing bank statement components: {', '.join(missing)}")
        found["revenue"] = (
            bank_interest_income - bank_interest_expense + bank_non_interest,
            "net interest income + non-interest income",
        )
        gross_profit_label = (
            "pre-provision operating profit"
            if bank_seen["provisions"]
            else "operating profit (provisions not separately reported)"
        )
        if bank_operating_profit_derived:
            gross_profit_label += " derived from pre-tax profit"
        found["gross_profit"] = (
            bank_operating_profit - bank_provision_effect,
            gross_profit_label,
        )
    return found, sources


def extract_workbook(
    content: bytes,
    expected_ticker: str | None = None,
    expected_end: dt.date | None = None,
) -> dict[str, Any]:
    if openpyxl is None:
        raise RuntimeError("openpyxl is required")
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
    meta = {}
    for row in (
        wb["1000000"].iter_rows(values_only=True) if "1000000" in wb.sheetnames else []
    ):
        if row and len(row) > 1:
            k = norm(row[0])
            meta[k] = row[1]
    ticker = (
        str(
            meta_value(meta, "kode entitas", "kode emiten", "kode emiten / stock code")
            or ""
        )
        .strip()
        .upper()
        .replace(".JK", "")
    )
    currency = str(
        meta_value(meta, "mata uang pelaporan", prefixes=("mata uang pelaporan",)) or ""
    ).strip()
    currency = currency.split("/")[-1].strip().upper()
    raw_end = meta_value(
        meta,
        "tanggal akhir periode berjalan",
        "tanggal akhir periode",
        "period end date",
    )
    raw_start = meta_value(
        meta,
        "tanggal awal periode berjalan",
        "tanggal mulai periode",
        "period start date",
    )
    end = parse_date(raw_end)
    if expected_ticker and ticker != expected_ticker.upper().replace(".JK", ""):
        raise ValueError(f"workbook ticker {ticker!r} != {expected_ticker}")
    if expected_end and end != expected_end:
        raise ValueError(f"period end {end} != {expected_end}")
    start = parse_date(raw_start)
    if start is None:
        raise ValueError(f"missing or unrecognized period start: {raw_start!r}")
    if (start.month, start.day) != (1, 1):
        raise ValueError(f"period starts {start}, not Jan 1")
    if end and start.year != end.year:
        raise ValueError(f"period start year {start.year} != end year {end.year}")
    mult = factor_value(
        meta_value(meta, "pembulatan", "rounding", prefixes=("pembulatan",))
    )
    classification = " ".join(
        norm(meta.get(label)) for label in ("subsektor", "industri", "subindustri")
    )
    is_bank = "bank" in classification
    candidates = []
    for name in wb.sheetnames:
        if not str(name).isdigit() or str(name)[:2] not in {
            "13",
            "23",
            "33",
            "43",
            "83",
        }:
            continue
        found, sources = extract_statement(wb[name], mult, is_bank)
        if "net_income" in found:
            candidates.append((str(name), found, sources))
    if not ticker or not currency or not end:
        raise ValueError("missing essential metadata")
    if not candidates:
        raise ValueError("missing exact parent net income")
    signatures = {
        tuple(
            candidate[1].get(key, (None, None))[0]
            for key in ("revenue", "gross_profit", "net_income")
        )
        for candidate in candidates
    }
    if len(signatures) != 1:
        sheets = ", ".join(candidate[0] for candidate in candidates)
        raise ValueError(f"ambiguous populated statement sheets: {sheets}")
    statement_sheet, found, sources = candidates[0]
    scope = str(
        meta_value(
            meta,
            "apakah merupakan laporan keuangan satu entitas atau suatu kelompok entitas",
            "scope",
        )
        or ""
    ).strip()
    return {
        "ticker": ticker,
        "currency": currency,
        "is_bank": is_bank,
        "scope": scope,
        "statement_sheet": statement_sheet,
        "period_end": end,
        "revenue": found.get("revenue", (None, None))[0],
        "gross_profit": found.get("gross_profit", (None, None))[0],
        "net_income": found["net_income"][0],
        "revenue_label": found.get("revenue", (None, None))[1],
        "gross_profit_label": found.get("gross_profit", (None, None))[1],
        "net_income_label": found["net_income"][1],
        "metric_sources": sources,
        "rounding_factor": mult,
    }


def standalone(current, previous):
    if previous is None:
        return None
    out = {}
    for k in ("revenue", "gross_profit", "net_income"):
        a, b = current.get(k), previous.get(k)
        out[k] = a - b if a is not None and b is not None else None
    return out


def request_timeout(timeout, deadline):
    if deadline is None:
        return timeout
    remaining = deadline - time.monotonic()
    if remaining <= 0:
        raise TimeoutError("network time budget exceeded")
    return max(0.1, min(timeout, remaining))


def retry_sleep(seconds, deadline):
    if deadline is not None:
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            raise TimeoutError("network time budget exceeded")
        seconds = min(seconds, remaining)
    time.sleep(seconds)


def discover_all(session, year, period, timeout=20, deadline=None):
    all_rows = []
    expected = None
    start = 0
    page = 1000
    seen_pages = set()
    while True:
        params = {
            "indexFrom": start,
            "pageSize": page,
            "year": year,
            "reportType": "rdf",
            "periode": period,
            "SortColumn": "File_Modified",
            "SortOrder": "desc",
        }
        r = None
        for attempt in range(3):
            try:
                r = session.get(
                    IDX_API,
                    params=params,
                    impersonate=IMPERSONATE,
                    timeout=request_timeout(timeout, deadline),
                )
                if r.status_code not in RETRY or attempt == 2:
                    break
            except Exception:
                if attempt == 2:
                    raise
            retry_sleep(1 + attempt * 2, deadline)
        if r.status_code != 200:
            raise RuntimeError(f"discovery {year}/{period} HTTP {r.status_code}")
        p = r.json()
        rows = p.get("Results") or []
        count = int(p.get("ResultCount", 0) or 0)
        # IDX occasionally emits ResultCount=0 while returning a real first page;
        # retry that page rather than accepting an unvalidated universe.
        if count == 0 and rows:
            r = session.get(
                IDX_API,
                params=params,
                impersonate=IMPERSONATE,
                timeout=request_timeout(timeout, deadline),
            )
            if r.status_code != 200:
                raise RuntimeError(
                    f"discovery retry {year}/{period} HTTP {r.status_code}"
                )
            p = r.json()
            rows = p.get("Results") or []
            count = int(p.get("ResultCount", 0) or 0)
        if count == 0 and rows:
            raise RuntimeError(f"unvalidated ResultCount for {year}/{period}")
        if expected is None:
            expected = count
        elif count != expected:
            raise RuntimeError(
                f"ResultCount changed for {year}/{period}: {expected} -> {count}"
            )
        if not expected:
            break
        marker = hash(json.dumps(rows, sort_keys=True, default=str))
        if marker in seen_pages:
            raise RuntimeError(f"repeated page for {year}/{period}")
        seen_pages.add(marker)
        all_rows.extend(rows)
        start += len(rows)
        if len(seen_pages) > 100:
            raise RuntimeError(f"pagination limit for {year}/{period}")
        if len(all_rows) >= expected:
            break
    if expected is not None and len(all_rows) != expected:
        raise RuntimeError(
            f"count mismatch {year}/{period}: {len(all_rows)}/{expected}"
        )
    return all_rows


def xlsx_attachment(report):
    if not report:
        return None
    a = [
        x
        for x in (report.get("Attachments") or [])
        if str(x.get("File_Name", "")).lower().endswith(".xlsx") and x.get("File_Path")
    ]
    a.sort(
        key=lambda x: (
            "financialstatement" not in str(x.get("File_Name", "")).lower(),
            str(x.get("File_Name", "")),
        )
    )
    return a[0] if a else None


def inline_attachment(report):
    if not report:
        return None
    return next(
        (
            attachment
            for attachment in (report.get("Attachments") or [])
            if str(attachment.get("File_Name", "")).lower() == "inlinexbrl.zip"
            and attachment.get("File_Path")
        ),
        None,
    )


def attachment_url(attachment):
    return urljoin(IDX_BASE + "/", str(attachment["File_Path"]))


def is_xlsx(content):
    if not content or not zipfile.is_zipfile(BytesIO(content)):
        return False
    with zipfile.ZipFile(BytesIO(content)) as archive:
        names = set(archive.namelist())
    return "[Content_Types].xml" in names and "xl/workbook.xml" in names


def is_inline_zip(content):
    if not content or not zipfile.is_zipfile(BytesIO(content)):
        return False
    with zipfile.ZipFile(BytesIO(content)) as archive:
        names = set(archive.namelist())
    return "1000000.html" in names and any(
        Path(name).stem.isdigit()
        and Path(name).stem[:2] in {"13", "23", "33", "43", "83"}
        for name in names
    )


def inline_cell_value(cell):
    for element in cell.iter():
        if element.tag.rsplit("}", 1)[-1] != "nonFraction":
            continue
        text = "".join(element.itertext()).strip()
        if not text or text == "-":
            return 0
        negative = text.startswith("(") and text.endswith(")")
        text = text.strip("() ")
        format_name = str(element.get("format", "")).lower()
        if "numdotcomma" in format_name:
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
        value = float(text)
        if element.get("sign") == "-" or negative:
            value = -abs(value)
        return int(value) if value.is_integer() else value
    text = " ".join("".join(cell.itertext()).split())
    return text or None


def xlsx_from_inline_zip(content):
    if openpyxl is None:
        raise RuntimeError("openpyxl is required")
    if not is_inline_zip(content):
        raise ValueError("not a supported inline XBRL archive")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    with zipfile.ZipFile(BytesIO(content)) as archive:
        for name in sorted(archive.namelist()):
            sheet_name = Path(name).stem
            if sheet_name != "1000000" and (
                not sheet_name.isdigit()
                or sheet_name[:2] not in {"13", "23", "33", "43", "83"}
            ):
                continue
            root = ElementTree.fromstring(archive.read(name))
            sheet = wb.create_sheet(sheet_name)
            for table_row in root.iter("{http://www.w3.org/1999/xhtml}tr"):
                cells = table_row.findall("{http://www.w3.org/1999/xhtml}td")
                if len(cells) < 2:
                    continue
                label = " ".join("".join(cells[0].itertext()).split())
                if label:
                    sheet.append([label, inline_cell_value(cells[1])])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def canonicalize_rows(rows):
    out = []
    for row in rows:
        out.append(
            {
                k: str(row.get(k, "") if row.get(k, "") is not None else "")
                for k in CSV_HEADER
            }
        )
    return sorted(out, key=lambda x: (x["date"], x["symbol"], x["period"]))


def read_rows(path):
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def atomic_write(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        shutil.copy2(path, path.with_suffix(path.suffix + ".bak"))
    fd, tmp = tempfile.mkstemp(prefix=path.name + ".", dir=path.parent)
    os.close(fd)
    try:
        if path.exists():
            os.chmod(tmp, path.stat().st_mode)
        with open(tmp, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=CSV_HEADER, lineterminator="\n")
            w.writeheader()
            w.writerows(canonicalize_rows(rows))
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


def atomic_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=path.name + ".", dir=path.parent)
    os.close(fd)
    try:
        if path.exists():
            os.chmod(tmp, path.stat().st_mode)
        with open(tmp, "w", encoding="utf-8") as f:
            f.write(json.dumps(value, indent=2, sort_keys=True) + "\n")
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


def fingerprint(report):
    return hashlib.sha256(
        json.dumps(report, sort_keys=True, default=str).encode()
    ).hexdigest()


def cache_path(cache, year, period, ticker, fp):
    return cache / str(year) / period / f"{ticker}-{fp[:16]}.xlsx"


def fetch(session, url, timeout=120, deadline=None):
    last = None
    for i in range(3):
        try:
            last = session.get(
                url, impersonate=IMPERSONATE, timeout=request_timeout(timeout, deadline)
            )
            if last.status_code not in RETRY or i == 2:
                return last
        except Exception:
            if i == 2:
                raise
        retry_sleep(1 + i, deadline)
    return last


def cached_xlsx(session, attachment, path, timeout=120, deadline=None):
    if path.exists():
        content = path.read_bytes()
        if is_xlsx(content):
            return content
    r = fetch(session, attachment_url(attachment), timeout, deadline)
    if r.status_code != 200 or not is_xlsx(r.content):
        raise RuntimeError(f"download HTTP {r.status_code} or not xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=path.name + ".", dir=path.parent)
    os.close(fd)
    try:
        Path(tmp).write_bytes(r.content)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)
    return r.content


def cached_inline_zip(session, attachment, path, timeout=120, deadline=None):
    if path.exists():
        content = path.read_bytes()
        if is_inline_zip(content):
            return content
    r = fetch(session, attachment_url(attachment), timeout, deadline)
    if r.status_code != 200 or not is_inline_zip(r.content):
        raise RuntimeError(f"inline XBRL download HTTP {r.status_code} or invalid zip")
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=path.name + ".", dir=path.parent)
    os.close(fd)
    try:
        Path(tmp).write_bytes(r.content)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)
    return r.content


def report_workbook(session, report, cache_dir, year, period, ticker, fp, deadline):
    attachment = xlsx_attachment(report)
    if not attachment:
        raise RuntimeError("missing xlsx attachment")
    path = cache_path(cache_dir, year, period, ticker, fp)
    try:
        return cached_xlsx(session, attachment, path, 120, deadline), attachment, path
    except RuntimeError as xlsx_error:
        fallback = inline_attachment(report)
        if not fallback:
            raise
        inline_path = path.with_suffix(".inline.zip")
        try:
            content = cached_inline_zip(
                session, fallback, inline_path, 120, deadline
            )
            return xlsx_from_inline_zip(content), fallback, inline_path
        except Exception as inline_error:
            raise RuntimeError(
                f"{xlsx_error}; inline XBRL fallback failed: {inline_error}"
            ) from inline_error


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument(
        "--symbols", help="comma-separated existing CSV tickers; default: all"
    )
    ap.add_argument(
        "--max-items", type=int, default=100, help="maximum queued filings to process"
    )
    ap.add_argument(
        "--max-seconds", type=int, default=600, help="overall network time budget"
    )
    ap.add_argument(
        "--period",
        choices=PERIODS,
        help="sync only this period while still retrieving its cumulative dependency",
    )
    ap.add_argument(
        "--dry-run", action="store_true", help="extract and validate without writing"
    )
    ap.add_argument("--state-dir", type=Path, default=TASK / "state")
    ap.add_argument("--csv-dir", type=Path, default=CSV_DIR)
    ap.add_argument("--cache-dir", type=Path, default=TASK / "cache")
    ap.add_argument(
        "--timeout",
        type=int,
        default=20,
        help="per-request timeout within the time budget",
    )
    a = ap.parse_args(argv)
    if a.max_items < 0:
        ap.error("--max-items must be non-negative")
    if a.max_seconds <= 0:
        ap.error("--max-seconds must be positive")
    if a.timeout <= 0:
        ap.error("--timeout must be positive")

    deadline = time.monotonic() + a.max_seconds
    periods_to_discover = discovery_periods(a.period)
    today = dt.date.today()
    years = (today.year, today.year - 1)
    managed_symbols = {p.stem.upper() for p in a.csv_dir.glob("*.csv")}
    symbols = managed_symbols
    if a.symbols:
        requested = {
            x.strip().upper().removesuffix(".JK")
            for x in a.symbols.split(",")
            if x.strip()
        }
        unknown = requested - managed_symbols
        if unknown:
            print(f"no existing CSV for: {', '.join(sorted(unknown))}", file=sys.stderr)
            return 2
        symbols = requested
    if not symbols:
        print("no existing CSV tickers", file=sys.stderr)
        return 2
    if requests is None:
        print("curl_cffi is required", file=sys.stderr)
        return 2

    a.state_dir.mkdir(parents=True, exist_ok=True)
    a.cache_dir.mkdir(parents=True, exist_ok=True)
    state_path = a.state_dir / "state.json"
    queue_path = a.state_dir / "queue.json"
    prov_path = a.state_dir / "provenance.json"
    lock_file = (a.state_dir / ".lock").open("w")
    try:
        try:
            import fcntl

            fcntl.flock(lock_file, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            print("another sync is running", file=sys.stderr)
            return 3
        except ImportError:
            print("advisory file locking is unavailable", file=sys.stderr)
            return 3

        first_run = not state_path.exists()
        try:
            state = (
                json.loads(state_path.read_text(encoding="utf-8"))
                if state_path.exists()
                else {
                    "fingerprints": {},
                    "failures": {},
                    "extractor_version": EXTRACTOR_VERSION,
                }
            )
            queue = (
                json.loads(queue_path.read_text(encoding="utf-8"))
                if queue_path.exists()
                else []
            )
            prov = (
                json.loads(prov_path.read_text(encoding="utf-8"))
                if prov_path.exists()
                else {}
            )
        except (OSError, json.JSONDecodeError) as e:
            print(f"failed to read sync state: {e}", file=sys.stderr)
            return 2
        if (
            not isinstance(state, dict)
            or not isinstance(queue, list)
            or not isinstance(prov, dict)
        ):
            print("sync state has an invalid structure", file=sys.stderr)
            return 2
        state.setdefault("fingerprints", {})
        state.setdefault("failures", {})
        if not isinstance(state["fingerprints"], dict) or not isinstance(
            state["failures"], dict
        ):
            print("sync state has invalid fingerprints or failures", file=sys.stderr)
            return 2
        if state.get("extractor_version") != EXTRACTOR_VERSION:
            state["extractor_version"] = EXTRACTOR_VERSION
            state["fingerprints"] = {}
            state["failures"] = {}

        global IMPERSONATE
        session = requests.Session()
        warmed = False
        for imp in ("chrome124", "safari15_5"):
            try:
                r = session.get(
                    IDX_HOME,
                    impersonate=imp,
                    timeout=request_timeout(a.timeout, deadline),
                )
                if r.status_code < 400:
                    IMPERSONATE = imp
                    warmed = True
                    break
            except Exception:
                pass
        if not warmed:
            print("IDX homepage warm-up failed", file=sys.stderr)
            return 2

        reports = {}
        try:
            for y in years:
                for p in periods_to_discover:
                    if period_date(y, p) > today:
                        continue
                    for report in discover_all(session, y, p, a.timeout, deadline):
                        t = (
                            str(report.get("KodeEmiten", ""))
                            .strip()
                            .upper()
                            .removesuffix(".JK")
                        )
                        if t not in symbols or not xlsx_attachment(report):
                            continue
                        report_key = (t, y, p)
                        old = reports.get(report_key)
                        if old is None or str(report.get("File_Modified", "")) > str(
                            old.get("File_Modified", "")
                        ):
                            reports[report_key] = report
        except Exception as e:
            print(f"discovery failed: {e}", file=sys.stderr)
            return 2

        active_keys = {f"{t}:{y}:{p}" for t, y, p in reports}
        valid_queue = []
        for item in queue:
            if not isinstance(item, dict):
                print("sync queue contains an invalid item", file=sys.stderr)
                return 2
            try:
                t = str(item["ticker"])
                y = int(item["year"])
                p = str(item["period"])
                key = str(item["key"])
            except (KeyError, TypeError, ValueError):
                print("sync queue contains an invalid item", file=sys.stderr)
                return 2
            if p not in PERIODS or key != f"{t}:{y}:{p}":
                print(f"sync queue contains an invalid key: {key}", file=sys.stderr)
                return 2
            if (
                (a.period is None or p == a.period)
                and t in managed_symbols
                and (t not in symbols or key in active_keys)
            ):
                valid_queue.append(item)
        queue = valid_queue

        # A first run adopts existing rows. Extractor upgrades deliberately
        # invalidate fingerprints so current/prior rows are recalculated.
        queued_by_key = {q["key"]: q for q in queue}
        csv_rows = {t: read_rows(a.csv_dir / f"{t}.csv") for t in symbols}
        for (t, y, p), report in reports.items():
            if a.period is not None and p != a.period:
                continue
            key = f"{t}:{y}:{p}"
            fp = fingerprint(report)
            date_s = str(period_date(y, p))
            has = any(
                r.get("date") == date_s
                and r.get("symbol") == t
                and r.get("period") == PERIOD_NAME[p]
                for r in csv_rows[t]
            )
            if first_run and has:
                state["fingerprints"][key] = fp
                continue
            dependency = PREVIOUS[p]
            dependency_report = reports.get((t, y, dependency)) if dependency else None
            dependency_changed = (
                dependency is not None
                and (
                    dependency_report is None
                    or state["fingerprints"].get(f"{t}:{y}:{dependency}")
                    != fingerprint(dependency_report)
                )
            )
            if state["fingerprints"].get(key) != fp or not has or dependency_changed:
                item = {
                    "key": key,
                    "ticker": t,
                    "year": y,
                    "period": p,
                    "fingerprint": fp,
                }
                if key in queued_by_key:
                    queued_by_key[key].update(item)
                else:
                    queue.append(item)
                    queued_by_key[key] = item

        if a.period is None:
            changed_keys = {
                q["key"]
                for q in queue
                if q["key"] in active_keys
                and state["fingerprints"].get(q["key"]) != q.get("fingerprint")
            }
            for (t, y, p), report in reports.items():
                pp = PREVIOUS[p]
                key = f"{t}:{y}:{p}"
                if pp and f"{t}:{y}:{pp}" in changed_keys and key not in queued_by_key:
                    item = {
                        "key": key,
                        "ticker": t,
                        "year": y,
                        "period": p,
                        "fingerprint": fingerprint(report),
                    }
                    queue.append(item)
                    queued_by_key[key] = item

        queue.sort(
            key=lambda q: (
                bool(state["failures"].get(q["key"])),
                -q["year"],
                PERIODS.index(q["period"]),
                q["ticker"],
            )
        )
        eligible = [q for q in queue if q["key"] in active_keys]
        processed = 0
        changed = 0
        errors = 0
        successful_keys = set()

        for item in eligible[: a.max_items]:
            if time.monotonic() >= deadline:
                break
            t, y, p = item["ticker"], item["year"], item["period"]
            key = item["key"]
            try:
                report = reports[(t, y, p)]
                source_fp = fingerprint(report)
                content, att, cp = report_workbook(
                    session, report, a.cache_dir, y, p, t, source_fp, deadline
                )
                cur = extract_workbook(
                    content, t, period_date(y, p)
                )
                prev = None
                pa = None
                pp = PREVIOUS[p]
                if pp:
                    prior = reports.get((t, y, pp))
                    if not prior:
                        raise RuntimeError(f"missing dependency {pp}")
                    pfp = fingerprint(prior)
                    previous_content, pa, _ = report_workbook(
                        session, prior, a.cache_dir, y, pp, t, pfp, deadline
                    )
                    prev = extract_workbook(
                        previous_content,
                        t,
                        period_date(y, pp),
                    )
                    if prev["currency"] != cur["currency"]:
                        raise RuntimeError(
                            f"currency mismatch {prev['currency']} -> {cur['currency']}"
                        )
                    if (
                        prev.get("scope")
                        and cur.get("scope")
                        and norm(prev["scope"]) != norm(cur["scope"])
                    ):
                        raise RuntimeError("consolidated/single-entity scope mismatch")
                vals = cur if p == "tw1" else standalone(cur, prev)
                if vals is None:
                    raise RuntimeError("cumulative dependency unavailable")
                row = {
                    "date": str(period_date(y, p)),
                    "symbol": t,
                    "reported_currency": cur["currency"],
                    "calendar_year": str(y),
                    "period": PERIOD_NAME[p],
                    "revenue": vals["revenue"],
                    "gross_profit": vals["gross_profit"],
                    "net_income": vals["net_income"],
                }
                path = a.csv_dir / f"{t}.csv"
                rows = read_rows(path)
                rows = [
                    r
                    for r in rows
                    if not (
                        r.get("date") == row["date"]
                        and r.get("symbol") == t
                        and r.get("period") == row["period"]
                    )
                ] + [row]
                for old in rows:
                    if (
                        old.get("symbol") == t
                        and old.get("calendar_year") == str(y)
                        and old.get("period") == PERIOD_NAME[p]
                        and old.get("date") != row["date"]
                    ):
                        raise RuntimeError(
                            f"existing noncanonical date {old['date']} for {t} {y} {PERIOD_NAME[p]}"
                        )
                warning = sorted(
                    {
                        r.get("reported_currency")
                        for r in rows
                        if r.get("reported_currency")
                        and r.get("reported_currency") != cur["currency"]
                    }
                )
                if not a.dry_run:
                    atomic_write(path, rows)
                    prov[key] = {
                        "fingerprint": source_fp,
                        "extractor_version": EXTRACTOR_VERSION,
                        "source_url": attachment_url(att),
                        "dependency_source_url": attachment_url(pa) if pa else None,
                        "xlsx": str(cp),
                        "labels": {
                            "revenue": cur["revenue_label"],
                            "gross_profit": cur["gross_profit_label"],
                            "net_income": cur["net_income_label"],
                        },
                        "metric_sources": cur["metric_sources"],
                        "statement_sheet": cur["statement_sheet"],
                        "is_bank": cur["is_bank"],
                        "currency": cur["currency"],
                        "scope": cur["scope"],
                        "currency_warning": warning,
                        "rounding_factor": cur["rounding_factor"],
                    }
                    state["fingerprints"][key] = source_fp
                    state["failures"].pop(key, None)
                successful_keys.add(key)
                processed += 1
                changed += 0 if a.dry_run else 1
            except Exception as e:
                if not a.dry_run:
                    state["failures"][key] = {
                        "error": str(e),
                        "at": dt.datetime.now(dt.timezone.utc).isoformat(),
                    }
                errors += 1

        remaining = [q for q in queue if a.dry_run or q["key"] not in successful_keys]
        if not a.dry_run:
            atomic_json(queue_path, remaining)
            atomic_json(state_path, state)
            atomic_json(prov_path, prov)
        print(
            json.dumps(
                {
                    "discovered": len(reports),
                    "queued": len(queue),
                    "processed": processed,
                    "changed": changed,
                    "errors": errors,
                    "remaining": len(remaining),
                    "dry_run": a.dry_run,
                },
                sort_keys=True,
            )
        )
        return 1 if errors else 0
    finally:
        lock_file.close()


if __name__ == "__main__":
    raise SystemExit(main())
